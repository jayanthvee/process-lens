"""Reading a Word document into columns and rows: the three shapes, and the plain refusal."""

from __future__ import annotations

import io
from pathlib import Path

import pytest
from docx import Document

from packages.documents.docx_table import (
    NO_READABLE_CONTENT,
    DocumentRefused,
    NotAWordDocument,
    read_docx,
)

FIXTURES = Path(__file__).resolve().parents[3] / "fixtures" / "documents"


def load(name: str) -> bytes:
    return (FIXTURES / name).read_bytes()


def test_a_plain_table_keeps_its_header_and_rows() -> None:
    columns, rows = read_docx(load("clean-table.docx"))

    assert columns == ["Name", "Email", "Phone"]
    assert rows == [
        {
            "Name": 'Ana "Ani" Ruiz, Jr.',
            "Email": "Ana@Example.test",
            "Phone": "+34 600 123 456",
        },
        {
            "Name": "Sam Lee",
            "Email": "sam@example.test",
            "Phone": "+1 (415) 555-0134",
        },
    ]


def test_a_header_split_across_two_rows_is_joined() -> None:
    columns, rows = read_docx(load("merged-header.docx"))

    # The merged "Contact" is repeated into both columns, and the second header row names them.
    assert columns == ["Contact Name", "Contact Phone", "Email"]
    assert rows[0] == {
        "Contact Name": "Ana Ruiz",
        "Contact Phone": "+34 600 123 456",
        "Email": "Ana@Example.test",
    }
    assert rows[1]["Contact Name"] == "Sam Lee"


def test_repeated_label_value_blocks_become_rows() -> None:
    columns, rows = read_docx(load("label-blocks.docx"))

    assert columns == ["Name", "Email", "Phone"]
    assert rows == [
        {
            "Name": "Ana Ruiz",
            "Email": "Ana@Example.test",
            "Phone": "+34 600 123 456",
        },
        {
            "Name": "Sam Lee",
            "Email": "sam@example.test",
            "Phone": "+1 (415) 555-0134",
        },
    ]


def test_a_bulleted_list_is_one_column_named_from_its_heading() -> None:
    columns, rows = read_docx(load("bulleted-list.docx"))

    assert columns == ["Candidates"]
    assert rows == [
        {"Candidates": "Ana Ruiz"},
        {"Candidates": "Sam Lee"},
        {"Candidates": "Priya Nair"},
    ]


def test_a_document_with_no_table_or_list_is_refused_in_one_sentence() -> None:
    with pytest.raises(DocumentRefused) as caught:
        read_docx(load("no-table.docx"))

    assert str(caught.value) == NO_READABLE_CONTENT
    assert str(caught.value) == "This document has no table or list ProcessLens can read."


def test_a_file_that_is_not_a_word_document_is_named_as_such() -> None:
    with pytest.raises(NotAWordDocument) as caught:
        read_docx(b"Name,Email\nAna,a@example.test\n")

    assert str(caught.value) == "This file is not a Word document."


def test_the_largest_table_with_a_header_is_the_one_read() -> None:
    document = Document()
    small = document.add_table(rows=2, cols=1)
    small.cell(0, 0).text = "Note"
    small.cell(1, 0).text = "just one line"
    large = document.add_table(rows=3, cols=2)
    for column, text in enumerate(["Name", "Email"]):
        large.cell(0, column).text = text
    for index, values in enumerate(
        [["Ana", "ana@example.test"], ["Sam", "sam@example.test"]], start=1
    ):
        for column, text in enumerate(values):
            large.cell(index, column).text = text
    buffer = io.BytesIO()
    document.save(buffer)

    columns, rows = read_docx(buffer.getvalue())

    assert columns == ["Name", "Email"]
    assert rows[0] == {"Name": "Ana", "Email": "ana@example.test"}


def test_a_list_with_no_heading_above_it_is_called_item() -> None:
    document = Document()
    document.add_paragraph("Ana Ruiz", style="List Bullet")
    document.add_paragraph("Sam Lee", style="List Bullet")
    buffer = io.BytesIO()
    document.save(buffer)

    columns, rows = read_docx(buffer.getvalue())

    assert columns == ["Item"]
    assert rows == [{"Item": "Ana Ruiz"}, {"Item": "Sam Lee"}]
