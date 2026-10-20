"""Read a Word document into columns and rows.

Three shapes are understood, tried in this order: a table, repeated ``Label: value`` blocks, and a
bulleted or numbered list. Anything else is refused in one sentence. The document is only read:
macros are never run, external links are never followed, and embedded objects are never opened.
"""

from __future__ import annotations

import io
import zipfile
from typing import Any

from docx import Document
from docx.opc.exceptions import PackageNotFoundError

NO_READABLE_CONTENT = "This document has no table or list ProcessLens can read."
NOT_A_WORD_DOCUMENT = "This file is not a Word document."


class DocumentRefused(Exception):
    """The document cannot become a table, and the sentence says why."""


class NotAWordDocument(DocumentRefused):
    """The upload is not a Word document, whatever its name says."""


class NoReadableContent(DocumentRefused):
    """The document is a Word document but has no table, label blocks or list."""


def _style_name(paragraph: Any) -> str:
    style = getattr(paragraph, "style", None)
    name = getattr(style, "name", "") if style is not None else ""
    return str(name)


def _is_heading(paragraph: Any) -> bool:
    name = _style_name(paragraph).lower()
    return name.startswith("heading") or name == "title"


def _is_list(paragraph: Any) -> bool:
    name = _style_name(paragraph).lower()
    if "list bullet" in name or "list number" in name:
        return True
    properties = getattr(paragraph._p, "pPr", None)
    return properties is not None and properties.numPr is not None


def _has_horizontal_merge(row: Any) -> bool:
    """True when two neighbouring grid columns are the same cell (a merged header)."""
    cells = row.cells
    keys = [id(getattr(cell, "_tc", cell)) for cell in cells]
    return any(left == right for left, right in zip(keys, keys[1:], strict=False))


def _joined_header(top: str, bottom: str) -> str:
    parts = [part.strip() for part in (top, bottom) if part.strip()]
    # A vertically merged cell repeats itself in both rows; say it once.
    if len(parts) == 2 and parts[0] == parts[1]:
        parts = parts[:1]
    return " ".join(parts)


def _best_table(document: Any) -> Any | None:
    """The largest table that has a header and at least one row under it."""
    candidates = []
    for table in document.tables:
        if len(table.rows) < 2:
            continue
        first = [cell.text for cell in table.rows[0].cells]
        if any(text.strip() for text in first):
            candidates.append(table)
    if not candidates:
        return None
    # ``max`` keeps the first of equal-sized tables, which is the "first table" in the document.
    return max(candidates, key=lambda table: len(table.rows))


def _read_table(table: Any) -> tuple[list[str], list[dict[str, str]]]:
    rows = table.rows
    first = [cell.text for cell in rows[0].cells]
    second = [cell.text for cell in rows[1].cells] if len(rows) > 1 else []
    two_row_header = (
        len(rows) >= 3
        and bool(second)
        and (
            _has_horizontal_merge(rows[0])
            or (any(not text.strip() for text in first) and all(text.strip() for text in second))
        )
    )

    if two_row_header:
        width = max(len(first), len(second))
        headers = [
            _joined_header(
                first[index] if index < len(first) else "",
                second[index] if index < len(second) else "",
            )
            for index in range(width)
        ]
        body = rows[2:]
    else:
        headers = list(first)
        body = rows[1:]

    result: list[dict[str, str]] = []
    for row in body:
        values = [cell.text for cell in row.cells]
        values += [""] * max(0, len(headers) - len(values))
        result.append(dict(zip(headers, values, strict=False)))
    return headers, result


def _label_blocks(document: Any) -> tuple[list[str], list[dict[str, str]]] | None:
    """Repeated blocks of ``Label: value`` lines, one row per block."""
    blocks: list[dict[str, str]] = []
    current: dict[str, str] = {}

    def flush() -> None:
        nonlocal current
        if current:
            blocks.append(current)
            current = {}

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text or _is_heading(paragraph) or _is_list(paragraph):
            flush()
            continue
        if ":" not in text:
            # Prose inside a block is not a label; it does not split the block either.
            continue
        label, value = text.split(":", 1)
        label, value = label.strip(), value.strip()
        if not label or not value:
            flush()
            continue
        current[label] = value
    flush()

    if not blocks:
        return None
    if len(blocks) == 1 and len(blocks[0]) < 2:
        return None

    columns: list[str] = []
    for block in blocks:
        for label in block:
            if label not in columns:
                columns.append(label)
    return columns, [{column: block.get(column, "") for column in columns} for block in blocks]


def _list_items(document: Any) -> tuple[list[str], str | None] | None:
    """The longest bulleted or numbered list, and the heading above it."""
    paragraphs = list(document.paragraphs)
    groups: list[tuple[int, list[str]]] = []
    current: list[str] = []
    start = 0

    for index, paragraph in enumerate(paragraphs):
        text = paragraph.text.strip()
        if _is_list(paragraph):
            if not text:
                continue
            if not current:
                start = index
            current.append(text)
        elif current:
            groups.append((start, current))
            current = []
    if current:
        groups.append((start, current))
    if not groups:
        return None

    start, items = max(groups, key=lambda group: len(group[1]))
    heading: str | None = None
    for paragraph in reversed(paragraphs[:start]):
        if _is_heading(paragraph) and paragraph.text.strip():
            heading = paragraph.text.strip()
            break
    return items, heading


def _open(content: bytes) -> Any:
    """Open the document, or say plainly that it is not a Word file.

    A .docx is a zip archive with ``word/document.xml``. Checking both before handing the bytes to
    python-docx keeps a renamed spreadsheet or a broken upload from reading as an empty document.
    The parser reads the package's own parts only; no macro, link or embedded object is executed.
    """
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            names = set(archive.namelist())
    except (zipfile.BadZipFile, OSError, RuntimeError) as exc:
        raise NotAWordDocument(NOT_A_WORD_DOCUMENT) from exc
    if "word/document.xml" not in names:
        raise NotAWordDocument(NOT_A_WORD_DOCUMENT)
    try:
        return Document(io.BytesIO(content))
    except (PackageNotFoundError, KeyError, OSError, SyntaxError, TypeError, ValueError) as exc:
        raise NotAWordDocument(NOT_A_WORD_DOCUMENT) from exc


def read_docx(content: bytes) -> tuple[list[str], list[dict[str, str]]]:
    """Return the document's columns and rows, or raise :class:`DocumentRefused`."""
    document = _open(content)

    table = _best_table(document)
    if table is not None:
        columns, rows = _read_table(table)
        if columns:
            return columns, rows

    blocks = _label_blocks(document)
    if blocks is not None:
        return blocks

    listed = _list_items(document)
    if listed is not None:
        items, heading = listed
        column = heading or "Item"
        return [column], [{column: item} for item in items]

    raise NoReadableContent(NO_READABLE_CONTENT)


__all__ = [
    "NO_READABLE_CONTENT",
    "NOT_A_WORD_DOCUMENT",
    "DocumentRefused",
    "NoReadableContent",
    "NotAWordDocument",
    "read_docx",
]
