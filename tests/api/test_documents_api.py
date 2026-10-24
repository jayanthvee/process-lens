"""Word documents over HTTP: the preview, the CSV download, and a batch made from one."""

from __future__ import annotations

import base64
import csv
import io
import uuid
from pathlib import Path

import psycopg
import pytest
from docx import Document
from fastapi.testclient import TestClient

from services.api import documents
from tests.db.seed import Ledger

REPO_ROOT = Path(__file__).resolve().parents[2]
FIXTURES = REPO_ROOT / "fixtures" / "documents"
DOCX_MEDIA = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
MEGABYTE = 1024 * 1024
REFUSAL = "This document has no table or list ProcessLens can read."


def load(name: str) -> bytes:
    return (FIXTURES / name).read_bytes()


def to_csv(
    client: TestClient,
    content: bytes,
    *,
    filename: str = "contacts.docx",
) -> object:
    return client.post(
        "/api/v1/documents/to-csv",
        json={
            "file_name": filename,
            "content_base64": base64.b64encode(content).decode("ascii"),
        },
    )


def table_document(rows: list[list[str]]) -> bytes:
    """A Word document with one table, built here so the rows are the test's own."""
    document = Document()
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    for index, values in enumerate(rows):
        for column, text in enumerate(values):
            table.cell(index, column).text = text
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def add_task(conn: psycopg.Connection, tenant_id: uuid.UUID) -> uuid.UUID:
    """One task with one recipe version; the upload only needs the workflow."""
    name = f"Word contacts {uuid.uuid4().hex[:6]}"
    with conn.cursor() as cur:
        cur.execute(
            "INSERT INTO workflows (tenant_id, name, destination_key, allowed_origins) "
            "VALUES (%s, %s, 'demo-crm', %s) RETURNING id",
            (tenant_id, name, ["http://localhost:5173"]),
        )
        workflow_id = cur.fetchone()[0]
        cur.execute(
            "INSERT INTO recipe_versions (tenant_id, workflow_id, version, recipe, change_reason) "
            "VALUES (%s, %s, 1, %s, 'compiled from recording')",
            (tenant_id, workflow_id, '{"format_version": 1, "meta": {"name": "Add contacts"}}'),
        )
        return workflow_id


def stored_rows(conn: psycopg.Connection, batch_id: str) -> list[dict]:
    with conn.cursor() as cur:
        cur.execute(
            "SELECT row_index, raw, clean, record_key, status FROM input_records "
            "WHERE batch_id = %s ORDER BY row_index",
            (batch_id,),
        )
        columns = ["row_index", "raw", "clean", "record_key", "status"]
        return [dict(zip(columns, row, strict=True)) for row in cur.fetchall()]


# --- the preview ------------------------------------------------------------------------------


def test_to_csv_answers_the_columns_the_first_rows_and_the_notes(connected: TestClient) -> None:
    response = to_csv(connected, load("clean-table.docx"))

    assert response.status_code == 200, response.text
    body = response.json()
    assert body["columns"] == ["Name", "Email", "Phone"]
    assert body["rows"] == [
        {
            "Name": 'Ana "Ani" Ruiz, Jr.',
            "Email": "ana@example.test",
            "Phone": "+34 600 123 456",
        },
        {
            "Name": "Sam Lee",
            "Email": "sam@example.test",
            "Phone": "+1 (415) 555-0134",
        },
    ]
    assert body["notes"] == []
    assert body["csv_token"]


def test_to_csv_returns_only_the_first_fifty_rows(connected: TestClient) -> None:
    rows = [["Name", "Email"]]
    rows += [[f"Person {index}", f"person{index}@example.test"] for index in range(60)]

    body = to_csv(connected, table_document(rows)).json()

    assert len(body["rows"]) == 50
    assert body["rows"][0] == {"Name": "Person 0", "Email": "person0@example.test"}
    assert body["rows"][-1] == {"Name": "Person 49", "Email": "person49@example.test"}


# --- the download -----------------------------------------------------------------------------


def test_the_download_is_utf8_with_a_bom_and_escapes_commas_and_quotes(
    connected: TestClient,
) -> None:
    token = to_csv(connected, load("clean-table.docx")).json()["csv_token"]

    response = connected.get(f"/api/v1/documents/to-csv/{token}.csv")

    assert response.status_code == 200, response.text
    assert response.content.startswith(b"\xef\xbb\xbf")
    text = response.content.decode("utf-8-sig")
    # A comma and a quote in a cell are written the way CSV requires.
    assert '"Ana ""Ani"" Ruiz, Jr."' in text
    parsed = list(csv.reader(io.StringIO(text)))
    assert parsed[0] == ["Name", "Email", "Phone"]
    assert parsed[1] == ['Ana "Ani" Ruiz, Jr.', "ana@example.test", "+34 600 123 456"]
    assert parsed[2] == ["Sam Lee", "sam@example.test", "+1 (415) 555-0134"]


def test_an_unknown_or_expired_token_is_a_plain_404(
    connected: TestClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    assert connected.get("/api/v1/documents/to-csv/nope.csv").status_code == 404

    monkeypatch.setattr(documents, "TOKEN_TTL_SECONDS", -1)
    token = to_csv(connected, load("clean-table.docx")).json()["csv_token"]
    assert connected.get(f"/api/v1/documents/to-csv/{token}.csv").status_code == 404


# --- refusals ---------------------------------------------------------------------------------


def test_a_file_over_five_megabytes_is_413(connected: TestClient) -> None:
    content = b"PK\x03\x04" + b"\0" * (5 * MEGABYTE)

    response = to_csv(connected, content, filename="too-big.docx")

    assert response.status_code == 413
    assert "5 MB" in response.json()["detail"]


def test_a_file_that_is_not_a_docx_renamed_as_one_is_422(connected: TestClient) -> None:
    response = to_csv(connected, b"Name,Email\nAna,a@example.test\n")

    assert response.status_code == 422
    assert response.json()["detail"] == "This file is not a Word document."


def test_a_document_with_no_table_or_list_is_422_in_one_sentence(connected: TestClient) -> None:
    response = to_csv(connected, load("no-table.docx"))

    assert response.status_code == 422
    assert response.json()["detail"] == REFUSAL


def test_only_a_docx_or_pdf_suffix_is_accepted(connected: TestClient) -> None:
    response = to_csv(connected, load("clean-table.docx"), filename="contacts.doc")

    assert response.status_code == 422
    assert response.json()["detail"] == "Only a .docx or .pdf file can be added here."


# --- the batch --------------------------------------------------------------------------------


def test_upload_batch_accepts_a_docx_and_stores_the_cleaned_rows(
    connected: TestClient, conn: psycopg.Connection, ledger: Ledger
) -> None:
    workflow_id = add_task(conn, ledger.tenant_id)

    response = connected.post(
        "/api/v1/batches",
        data={"workflow_id": str(workflow_id)},
        files={"file": ("contacts.docx", load("clean-table.docx"), DOCX_MEDIA)},
    )

    assert response.status_code == 201, response.text
    body = response.json()
    assert body["columns"] == ["Name", "Email", "Phone"]
    assert body["row_count"] == 2
    assert body["notes"] == []

    rows = stored_rows(conn, body["batch_id"])
    assert [row["row_index"] for row in rows] == [1, 2]
    assert rows[0]["raw"] == {
        "Name": 'Ana "Ani" Ruiz, Jr.',
        "Email": "ana@example.test",
        "Phone": "+34 600 123 456",
    }
    assert rows[0]["clean"] == rows[0]["raw"]
    assert rows[0]["record_key"] == "ana@example.test"
    assert {row["status"] for row in rows} == {"ready"}

    with conn.cursor() as cur:
        cur.execute("SELECT kind FROM input_batches WHERE id = %s", (body["batch_id"],))
        assert cur.fetchone()[0] == "csv"


def test_upload_batch_reports_the_cleaner_notes(
    connected: TestClient, conn: psycopg.Connection, ledger: Ledger
) -> None:
    workflow_id = add_task(conn, ledger.tenant_id)
    content = table_document(
        [
            ["Name", "Email"],
            ["Ana", "Ana@Example.test"],
            ["Ana", "ANA@example.test"],
            ["", ""],
            ["Sam", "sam@example.test"],
        ]
    )

    response = connected.post(
        "/api/v1/batches",
        data={"workflow_id": str(workflow_id)},
        files={"file": ("contacts.docx", content, DOCX_MEDIA)},
    )

    assert response.status_code == 201, response.text
    body = response.json()
    assert body["row_count"] == 2
    assert body["notes"] == ["Removed 1 empty row", "Removed 1 duplicate row"]
    assert [row["raw"]["Name"] for row in stored_rows(conn, body["batch_id"])] == ["Ana", "Sam"]


def test_upload_batch_refuses_a_docx_with_nothing_to_read(
    connected: TestClient, conn: psycopg.Connection, ledger: Ledger
) -> None:
    workflow_id = add_task(conn, ledger.tenant_id)

    response = connected.post(
        "/api/v1/batches",
        data={"workflow_id": str(workflow_id)},
        files={"file": ("contacts.docx", load("no-table.docx"), DOCX_MEDIA)},
    )

    assert response.status_code == 422
    assert response.json()["detail"] == REFUSAL


# --- PDFs (PL-091) ----------------------------------------------------------------------------

PDF_MEDIA = "application/pdf"


def test_to_csv_reads_a_pdf_of_invoices(connected: TestClient) -> None:
    response = to_csv(connected, load("invoices-one-per-page.pdf"), filename="invoices.pdf")

    assert response.status_code == 200, response.text
    body = response.json()
    assert body["columns"][:3] == ["Invoice Number", "Invoice Date", "Customer"]
    assert [row["Invoice Number"] for row in body["rows"]] == [
        "INV-2026-0101",
        "INV-2026-0102",
        "INV-2026-0103",
    ]
    assert body["csv_token"]


def test_to_csv_refuses_a_scanned_pdf_in_one_sentence(connected: TestClient) -> None:
    response = to_csv(connected, load("scanned.pdf"), filename="scan.pdf")

    assert response.status_code == 422
    assert "scanned" in response.json()["detail"]


def test_to_csv_refuses_a_docx_renamed_as_a_pdf(connected: TestClient) -> None:
    response = to_csv(connected, load("clean-table.docx"), filename="contacts.pdf")

    assert response.status_code == 422
    assert response.json()["detail"] == "This file is not a PDF."


def test_upload_batch_accepts_a_pdf_and_stores_one_row_per_invoice(
    connected: TestClient, conn: psycopg.Connection, ledger: Ledger
) -> None:
    workflow_id = add_task(conn, ledger.tenant_id)

    response = connected.post(
        "/api/v1/batches",
        data={"workflow_id": str(workflow_id)},
        files={"file": ("invoices.pdf", load("invoices-one-per-page.pdf"), PDF_MEDIA)},
    )

    assert response.status_code == 201, response.text
    body = response.json()
    assert body["row_count"] == 3
    assert "Unit Price" in body["columns"]

    rows = stored_rows(conn, body["batch_id"])
    assert [row["raw"]["Customer"] for row in rows] == [
        "Ridgeline Outfitters",
        "Harbourview Dental",
        "Pinewood Preschool: Main Campus",
    ]
    assert rows[1]["raw"]["Unit Price"] == "€1,299.00"

    with conn.cursor() as cur:
        cur.execute("SELECT kind FROM input_batches WHERE id = %s", (body["batch_id"],))
        assert cur.fetchone()[0] == "csv"


def test_upload_batch_refuses_a_pdf_with_nothing_to_read(
    connected: TestClient, conn: psycopg.Connection, ledger: Ledger
) -> None:
    workflow_id = add_task(conn, ledger.tenant_id)

    response = connected.post(
        "/api/v1/batches",
        data={"workflow_id": str(workflow_id)},
        files={"file": ("letter.pdf", load("letter.pdf"), PDF_MEDIA)},
    )

    assert response.status_code == 422
    assert (
        response.json()["detail"]
        == "This PDF has no labelled fields or table ProcessLens can read."
    )


def test_upload_batch_names_pdf_among_the_accepted_files(
    connected: TestClient, conn: psycopg.Connection, ledger: Ledger
) -> None:
    workflow_id = add_task(conn, ledger.tenant_id)

    response = connected.post(
        "/api/v1/batches",
        data={"workflow_id": str(workflow_id)},
        files={"file": ("invoices.doc", b"not used", "application/msword")},
    )

    assert response.status_code == 422
    assert response.json()["detail"] == (
        "Only a .csv, .tsv, .xlsx, .docx or .pdf file can be added here."
    )
